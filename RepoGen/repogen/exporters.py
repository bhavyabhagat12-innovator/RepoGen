from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

def clean(text: str) -> str:
    replacements = {
        "\u2011": "-",   # non-breaking hyphen
        "\u2010": "-",   # hyphen
        "\u2013": "-",   # en dash
        "\u2014": "-",   # em dash
        "\u2018": "'", "\u2019": "'",   # curly single quotes
        "\u201c": '"', "\u201d": '"',   # curly double quotes
        "\u00a0": " ",   # non-breaking space
        "\u202f": " ",   # narrow non-breaking space
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text

def to_docx(report: dict) -> bytes:
    doc = Document()
    doc.add_heading(report["title"], 0)
    for key, value in report["meta"].items():
        p = doc.add_paragraph()
        p.add_run(f"{key}: ").bold = True
        p.add_run(str(value))
    for s in report["sections"]:
        doc.add_heading(s["heading"], 1)
        for para in s["text"].split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(report: dict) -> bytes:
    buf = BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()

    story = [Paragraph(escape(report["title"]), styles["Title"]), Spacer(1, 12)]
    for key, value in report["meta"].items():
        story.append(Paragraph(f"<b>{escape(key)}:</b> {escape(str(value))}", styles["Normal"]))
    story.append(Spacer(1, 16))

    for s in report["sections"]:
        story.append(Paragraph(escape(clean(s["heading"])), styles["Heading2"]))
        for para in s["text"].split("\n\n"):
            if para.strip():
                story.append(Paragraph(escape(clean(para.strip())), styles["BodyText"]))
                story.append(Spacer(1, 8))

    pdf.build(story)
    return buf.getvalue()