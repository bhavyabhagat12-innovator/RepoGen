from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Image, Paragraph, SimpleDocTemplate, Spacer

from .header import Header

LOGO_WIDTH_INCH = 1.2


def clean(text: str) -> str:
    replacements = {
        "\u2011": "-", "\u2010": "-", "\u2013": "-", "\u2014": "-",
        "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"',
        "\u00a0": " ", "\u202f": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


# ---------- DOCX ----------

def _bottom_border(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    pPr.append(borders)


def _docx_header(doc, header: Header) -> None:
    last = None
    if header.logo_bytes:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(BytesIO(header.logo_bytes), width=Inches(LOGO_WIDTH_INCH))
        last = p
    for i, line in enumerate(header.lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.bold = i == 0
        run.font.size = Pt(16 if i == 0 else 11)
        last = p
    if last is not None:
        _bottom_border(last)
        doc.add_paragraph()


def to_docx(report: dict, header: Header | None = None) -> bytes:
    doc = Document()
    if header and not header.is_empty:
        _docx_header(doc, header)
    doc.add_heading(report["title"], 0)
    for key, value in report["meta"].items():
        p = doc.add_paragraph()
        p.add_run(f"{key}: ").bold = True
        p.add_run(str(value))
    for s in report["sections"]:
        doc.add_heading(s["heading"], 1)
        for para in s["text"].split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

def _pdf_header(header: Header, styles) -> list:
    items = []
    if header.logo_bytes:
        with PILImage.open(BytesIO(header.logo_bytes)) as im:
            w, h = im.size
        width = LOGO_WIDTH_INCH * 72
        img = Image(BytesIO(header.logo_bytes), width=width, height=width * h / w)
        img.hAlign = "CENTER"
        items += [img, Spacer(1, 6)]

    title_style = ParagraphStyle("HeaderTitle", parent=styles["Title"], fontSize=16, leading=20, spaceAfter=2)
    line_style = ParagraphStyle("HeaderLine", parent=styles["Normal"], alignment=TA_CENTER)
    for i, line in enumerate(header.lines):
        items.append(Paragraph(escape(clean(line)), title_style if i == 0 else line_style))

    items += [
        Spacer(1, 6),
        HRFlowable(width="100%", thickness=0.8, color=colors.grey),
        Spacer(1, 14),
    ]
    return items


def to_pdf(report: dict, header: Header | None = None) -> bytes:
    buf = BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()

    story = []
    if header and not header.is_empty:
        story += _pdf_header(header, styles)

    story += [Paragraph(escape(clean(report["title"])), styles["Title"]), Spacer(1, 12)]
    for key, value in report["meta"].items():
        story.append(
            Paragraph(f"<b>{escape(clean(key))}:</b> {escape(clean(str(value)))}", styles["Normal"])
        )
    story.append(Spacer(1, 16))

    for s in report["sections"]:
        story.append(Paragraph(escape(clean(s["heading"])), styles["Heading2"]))
        for para in s["text"].split("\n\n"):
            if para.strip():
                story.append(Paragraph(escape(clean(para.strip())), styles["BodyText"]))
                story.append(Spacer(1, 8))

    pdf.build(story)
    return buf.getvalue()